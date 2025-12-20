import os
import re
import unicodedata
from io import BytesIO
from uuid import uuid4
from datetime import datetime, date

import pandas as pd
from dotenv import load_dotenv
from flask import (
    Flask, render_template, request, redirect, url_for,
    flash, session, abort, send_file
)

from sqlalchemy import (
    create_engine, Column, Integer, String, Text,
    DateTime, ForeignKey, func
)
from sqlalchemy.orm import (
    scoped_session, sessionmaker, declarative_base, relationship
)
from sqlalchemy.dialects.postgresql import JSONB

from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

# ----------------------------
# Env
# ----------------------------
load_dotenv()

# ----------------------------
# Flask
# ----------------------------
app = Flask(__name__)
app.secret_key = os.getenv("SECRET_KEY", "dev-secret-key-change-me")
app.config["MAX_CONTENT_LENGTH"] = 20 * 1024 * 1024  # 20 MB

ALLOWED_EXTS = {"xlsx", "xlsm", "xls"}

# ----------------------------
# DB (PostgreSQL en Render)
# ----------------------------
DATABASE_URL = os.getenv("DATABASE_URL", "").strip()
if DATABASE_URL.startswith("postgres://"):
    DATABASE_URL = DATABASE_URL.replace("postgres://", "postgresql://", 1)

if not DATABASE_URL:
    raise RuntimeError("Falta DATABASE_URL (Render Environment Variables).")

DB_SSLMODE = os.getenv("DB_SSLMODE", "require")

engine = create_engine(
    DATABASE_URL,
    pool_pre_ping=True,
    connect_args={"sslmode": DB_SSLMODE},
)

db_session = scoped_session(sessionmaker(bind=engine, autoflush=False, autocommit=False))
Base = declarative_base()
Base.query = db_session.query_property()


class UploadBatch(Base):
    __tablename__ = "upload_batches"
    id = Column(Integer, primary_key=True)
    token = Column(String(64), unique=True, nullable=False, index=True)
    created_at = Column(DateTime(timezone=True), server_default=func.now(), nullable=False)

    ontracking_rows = relationship("OntrackingRow", back_populates="batch", cascade="all, delete-orphan")
    cardlog_rows = relationship("CardLogRow", back_populates="batch", cascade="all, delete-orphan")
    roommap_rows = relationship("RoomMapRow", back_populates="batch", cascade="all, delete-orphan")


class OntrackingRow(Base):
    __tablename__ = "ontracking_rows"
    id = Column(Integer, primary_key=True)
    batch_id = Column(Integer, ForeignKey("upload_batches.id"), nullable=False, index=True)

    modulo = Column(String(50))
    lugar = Column(String(80))
    habitacion = Column(String(80), index=True)
    empresa = Column(Text)
    ontracking_id = Column(String(80))
    cama = Column(String(20))
    inicio = Column(String(40))    # "YYYY-MM-DD"
    termino = Column(String(40))   # "YYYY-MM-DD"
    dia = Column(String(40))
    camas_ocupadas = Column(String(40))
    rut = Column(String(40), index=True)
    nombre = Column(Text)

    raw = Column(JSONB, nullable=False)
    batch = relationship("UploadBatch", back_populates="ontracking_rows")


class CardLogRow(Base):
    __tablename__ = "cardlog_rows"
    id = Column(Integer, primary_key=True)
    batch_id = Column(Integer, ForeignKey("upload_batches.id"), nullable=False, index=True)

    nro_tarjeta = Column(String(50), index=True)
    nro_habitacion = Column(String(80), index=True)
    habitacion = Column(String(80), index=True)  # nomenclatura log (ej: L1801-2)
    metodo_apertura_puerta = Column(Text)
    tipo_tarjeta = Column(String(80))
    fecha = Column(Text)  # string original
    dueno_codigo = Column(String(80))
    dueno_nombre = Column(Text)

    raw = Column(JSONB, nullable=False)
    batch = relationship("UploadBatch", back_populates="cardlog_rows")


class RoomMapRow(Base):
    __tablename__ = "roommap_rows"
    id = Column(Integer, primary_key=True)
    batch_id = Column(Integer, ForeignKey("upload_batches.id"), nullable=False, index=True)

    habitacion = Column(String(80), index=True)  # nomenclatura ontracking
    modulo = Column(String(80), index=True)
    piso = Column(String(20))
    hkeyplus = Column(String(80), index=True)    # nomenclatura log (ej: L1801-2)

    raw = Column(JSONB, nullable=False)
    batch = relationship("UploadBatch", back_populates="roommap_rows")


def init_db():
    Base.metadata.create_all(bind=engine)


@app.teardown_appcontext
def shutdown_session(exception=None):
    db_session.remove()


init_db()

# ----------------------------
# Helpers
# ----------------------------
def _normalize_col_name(name: str) -> str:
    name = str(name).strip()
    name = unicodedata.normalize("NFKD", name)
    name = "".join(ch for ch in name if not unicodedata.combining(ch))
    name = re.sub(r"[^A-Za-z0-9]+", "_", name)
    return name.strip("_").upper()


def normalize_columns(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    seen = {}
    cols = []
    for c in df.columns:
        base = _normalize_col_name(c)
        if base in seen:
            seen[base] += 1
            base = f"{base}_{seen[base]}"
        else:
            seen[base] = 1
        cols.append(base)
    df.columns = cols
    return df


def allowed_filename(filename: str) -> bool:
    if not filename or "." not in filename:
        return False
    ext = filename.rsplit(".", 1)[1].lower()
    return ext in ALLOWED_EXTS


def read_excel_upload(file_storage) -> pd.DataFrame:
    filename = file_storage.filename or ""
    if not allowed_filename(filename):
        raise ValueError(f"Archivo no permitido: {filename}")

    ext = filename.rsplit(".", 1)[1].lower()
    engine_name = "xlrd" if ext == "xls" else "openpyxl"

    raw = file_storage.read()
    if not raw:
        raise ValueError(f"Archivo vacío: {filename}")

    df = pd.read_excel(BytesIO(raw), dtype=str, engine=engine_name)
    df = normalize_columns(df).fillna("")
    return df


def rename_by_candidates(df: pd.DataFrame, candidates_map: dict[str, list[str]]) -> pd.DataFrame:
    df = df.copy()
    colset = set(df.columns)
    rename = {}
    for canonical, candidates in candidates_map.items():
        for cand in candidates:
            if cand in colset:
                rename[cand] = canonical
                break
    return df.rename(columns=rename)


def validate_required(df: pd.DataFrame, required: set, label: str) -> list[str]:
    missing = [c for c in sorted(required) if c not in df.columns]
    if missing:
        return [f"{label}: faltan columnas: {', '.join(missing)}"]
    return []


def parse_log_datetime(value: str):
    """
    Convierte strings del log a datetime.
    Soporta formatos con "a.m./p.m." y dayfirst.
    """
    if value is None:
        return None
    s = str(value).strip()
    if not s:
        return None
    s = s.replace("a.m.", "AM").replace("p.m.", "PM").replace("A.M.", "AM").replace("P.M.", "PM")
    try:
        dt = pd.to_datetime(s, errors="coerce", dayfirst=True)
        if pd.isna(dt):
            return None
        return dt.to_pydatetime()
    except Exception:
        return None


def clamp_int(x, default, lo, hi):
    try:
        v = int(x)
    except Exception:
        return default
    return max(lo, min(hi, v))


# ----------------------------
# Canonicalización
# ----------------------------
ONTRACKING_COLMAP = {
    "MODULO": ["MODULO", "MODU", "MOD"],
    "LUGAR": ["LUGAR", "LUGA", "LUG"],
    "HABITACION": ["HABITACION", "HABITACI", "HABITA", "HAB"],
    "EMPRESA": ["EMPRESA", "EMPRES"],
    "ID": ["ID"],
    "CAMA": ["CAMA", "CAM", "CAR"],
    "INICIO": ["INICIO", "INICI"],
    "TERMINO": ["TERMINO", "TERMIN", "TERM"],
    "DIA": ["DIA"],
    "CAMAS_OCUPADAS": ["CAMAS_OCUPADAS", "CAMAS_OCUPD", "CAMAS_OCUPDAS", "CAMAS_OCUP"],
    "RUT": ["RUT"],
    "NOMBRE": ["NOMBRE", "NOMBR"],
}

CARDLOG_COLMAP = {
    "NRO_TARJETA": ["NRO_TARJETA", "NRO_TARJET", "NRO_TARJ"],
    "NRO_HABITACION": ["NRO_HABITACION", "NRO_HABITAC", "NRO_HAB"],
    "HABITACION": ["HABITACION", "HABITACI", "HABITA"],
    "METODO_APERTURA_PUERTA": ["METODO_APERTURA_PUERTA", "METODO_APERTURA", "METODO"],
    "TIPO_DE_TARJETA": ["TIPO_DE_TARJETA", "TIPO_TARJETA", "TIPO"],
    "FECHA": ["FECHA"],
    "DUENO": ["DUENO", "DUEÑO", "OWNER"],
    "DUENO_2": ["DUENO_2", "DUEÑO_2"],
}

ROOMMAP_COLMAP = {
    "HABITACION": ["HABITACION", "HABITACI", "HABITA", "HAB"],
    "MODULO": ["MODULO", "MODU", "MOD"],
    "PISO": ["PISO"],
    "HKEYPLUS": ["HKEYPLUS", "HKEY_PLUS", "HKEY", "HKEYPLU"],
}


def canonicalize_ontracking(df: pd.DataFrame) -> pd.DataFrame:
    df = rename_by_candidates(df, ONTRACKING_COLMAP)
    for c in ("RUT", "HABITACION", "MODULO"):
        if c in df.columns:
            df[c] = df[c].astype(str).str.strip()

    for c in ("INICIO", "TERMINO"):
        if c in df.columns:
            dt = pd.to_datetime(df[c], errors="coerce", dayfirst=True)
            df[c] = dt.dt.date.astype(str).replace("NaT", "")
    return df


def canonicalize_cardlog(df: pd.DataFrame) -> pd.DataFrame:
    df = rename_by_candidates(df, CARDLOG_COLMAP)

    if "DUENO" in df.columns and "DUENO_2" in df.columns:
        df = df.rename(columns={"DUENO": "DUENO_CODIGO", "DUENO_2": "DUENO_NOMBRE"})
    elif "DUENO" in df.columns and "DUENO_CODIGO" not in df.columns:
        df = df.rename(columns={"DUENO": "DUENO_CODIGO"})

    for c in ("NRO_TARJETA", "NRO_HABITACION", "HABITACION"):
        if c in df.columns:
            df[c] = df[c].astype(str).str.strip()
    return df


def canonicalize_roommap(df: pd.DataFrame) -> pd.DataFrame:
    df = rename_by_candidates(df, ROOMMAP_COLMAP)
    for c in ("HABITACION", "MODULO", "HKEYPLUS", "PISO"):
        if c in df.columns:
            df[c] = df[c].astype(str).str.strip()
    return df


# ----------------------------
# Plantillas Excel
# ----------------------------
TEMPLATE_COLUMNS = {
    "ontracking": [
        "MODULO", "LUGAR", "HABITACION", "EMPRESA", "ID", "CAMA",
        "INICIO", "TERMINO", "DIA", "CAMAS_OCUPADAS", "RUT", "NOMBRE"
    ],
    "log_tarjetas": [
        "NRO_TARJETA", "NRO_HABITACION", "HABITACION",
        "METODO_APERTURA_PUERTA", "TIPO_DE_TARJETA", "FECHA",
        "DUENO_CODIGO", "DUENO_NOMBRE"
    ],
    "mapa_habitaciones": [
        "HABITACION", "MODULO", "PISO", "HKEYPLUS"
    ],
}

TEMPLATE_SHEETS = {
    "ontracking": "Ontracking",
    "log_tarjetas": "LogTarjetas",
    "mapa_habitaciones": "MapaHabitaciones",
}


def build_template_xlsx(template_key: str) -> BytesIO:
    cols = TEMPLATE_COLUMNS[template_key]
    sheet = TEMPLATE_SHEETS[template_key]
    df = pd.DataFrame(columns=cols)

    bio = BytesIO()
    with pd.ExcelWriter(bio, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name=sheet)
        ws = writer.book[sheet]
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        header_font = Font(bold=True)
        header_alignment = Alignment(horizontal="center", vertical="center")
        for cell in ws[1]:
            cell.font = header_font
            cell.alignment = header_alignment

        for i, col_name in enumerate(cols, start=1):
            width = max(14, min(42, len(col_name) + 2))
            ws.column_dimensions[get_column_letter(i)].width = width

    bio.seek(0)
    return bio


# ----------------------------
# Context
# ----------------------------
@app.context_processor
def inject_now():
    return {"now": datetime.now()}


# ----------------------------
# Routes
# ----------------------------
@app.get("/healthz")
def healthz():
    return {"status": "ok"}


@app.get("/")
def index():
    return redirect(url_for("importar"))


@app.get("/plantillas")
def plantillas_redirect():
    return redirect(url_for("importar"))


@app.get("/plantilla/<template_key>.xlsx")
def descargar_plantilla(template_key: str):
    if template_key not in TEMPLATE_COLUMNS:
        abort(404)
    bio = build_template_xlsx(template_key)
    return send_file(
        bio,
        as_attachment=True,
        download_name=f"plantilla_{template_key}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


@app.route("/importar", methods=["GET", "POST"])
def importar():
    if request.method == "GET":
        return render_template("importar.html", last_token=session.get("last_token"))

    token_in = (request.form.get("batch_token") or "").strip()

    f_ocup = request.files.get("excel_ocupacion")
    f_log = request.files.get("excel_log_tarjetas")
    f_map = request.files.get("excel_mapa_habitaciones")

    has_ocup = bool(f_ocup and (f_ocup.filename or "").strip())
    has_log = bool(f_log and (f_log.filename or "").strip())
    has_map = bool(f_map and (f_map.filename or "").strip())

    if not (has_ocup or has_log or has_map):
        flash("Debes subir al menos 1 archivo.", "danger")
        return redirect(url_for("importar"))

    try:
        # Batch
        if token_in:
            batch = db_session.query(UploadBatch).filter_by(token=token_in).first()
            if not batch:
                flash("Token no válido.", "danger")
                return redirect(url_for("importar"))
            token = batch.token
        else:
            token = uuid4().hex
            batch = UploadBatch(token=token)
            db_session.add(batch)
            db_session.flush()

        saved_sets = []

        # Ontracking
        if has_ocup:
            df = canonicalize_ontracking(read_excel_upload(f_ocup))
            errors = validate_required(df, {"MODULO", "LUGAR", "HABITACION", "RUT", "NOMBRE"}, "Ontracking")
            if errors:
                for e in errors:
                    flash(e, "danger")
                db_session.rollback()
                return redirect(url_for("importar"))

            db_session.query(OntrackingRow).filter_by(batch_id=batch.id).delete(synchronize_session=False)

            rows = []
            for r in df.to_dict(orient="records"):
                rows.append({
                    "batch_id": batch.id,
                    "modulo": r.get("MODULO", ""),
                    "lugar": r.get("LUGAR", ""),
                    "habitacion": r.get("HABITACION", ""),
                    "empresa": r.get("EMPRESA", ""),
                    "ontracking_id": r.get("ID", ""),
                    "cama": r.get("CAMA", ""),
                    "inicio": r.get("INICIO", ""),
                    "termino": r.get("TERMINO", ""),
                    "dia": r.get("DIA", ""),
                    "camas_ocupadas": r.get("CAMAS_OCUPADAS", ""),
                    "rut": r.get("RUT", ""),
                    "nombre": r.get("NOMBRE", ""),
                    "raw": r,
                })
            if rows:
                db_session.bulk_insert_mappings(OntrackingRow, rows)
            saved_sets.append("Ontracking")

        # Log Tarjetas
        if has_log:
            df = canonicalize_cardlog(read_excel_upload(f_log))
            errors = validate_required(df, {"NRO_TARJETA", "NRO_HABITACION", "HABITACION", "FECHA"}, "Log Tarjetas")
            if errors:
                for e in errors:
                    flash(e, "danger")
                db_session.rollback()
                return redirect(url_for("importar"))

            db_session.query(CardLogRow).filter_by(batch_id=batch.id).delete(synchronize_session=False)

            rows = []
            for r in df.to_dict(orient="records"):
                rows.append({
                    "batch_id": batch.id,
                    "nro_tarjeta": r.get("NRO_TARJETA", ""),
                    "nro_habitacion": r.get("NRO_HABITACION", ""),
                    "habitacion": r.get("HABITACION", ""),
                    "metodo_apertura_puerta": r.get("METODO_APERTURA_PUERTA", ""),
                    "tipo_tarjeta": r.get("TIPO_DE_TARJETA", ""),
                    "fecha": r.get("FECHA", ""),
                    "dueno_codigo": r.get("DUENO_CODIGO", ""),
                    "dueno_nombre": r.get("DUENO_NOMBRE", ""),
                    "raw": r,
                })
            if rows:
                db_session.bulk_insert_mappings(CardLogRow, rows)
            saved_sets.append("Log Tarjetas")

        # Mapa
        if has_map:
            df = canonicalize_roommap(read_excel_upload(f_map))
            errors = validate_required(df, {"HABITACION", "MODULO", "PISO", "HKEYPLUS"}, "Mapa")
            if errors:
                for e in errors:
                    flash(e, "danger")
                db_session.rollback()
                return redirect(url_for("importar"))

            db_session.query(RoomMapRow).filter_by(batch_id=batch.id).delete(synchronize_session=False)

            rows = []
            for r in df.to_dict(orient="records"):
                rows.append({
                    "batch_id": batch.id,
                    "habitacion": r.get("HABITACION", ""),
                    "modulo": r.get("MODULO", ""),
                    "piso": r.get("PISO", ""),
                    "hkeyplus": r.get("HKEYPLUS", ""),
                    "raw": r,
                })
            if rows:
                db_session.bulk_insert_mappings(RoomMapRow, rows)
            saved_sets.append("Mapa")

        db_session.commit()
        session["last_token"] = token
        flash(f"Guardado: {', '.join(saved_sets)} · Token: {token}", "success")
        return redirect(url_for("preview", token=token))

    except Exception as ex:
        db_session.rollback()
        flash(f"Error: {ex}", "danger")
        return redirect(url_for("importar"))


# ----------------------------
# PREVIEW con paginación + tabs
# ----------------------------
@app.get("/preview/<token>")
def preview(token: str):
    batch = db_session.query(UploadBatch).filter_by(token=token).first()
    if not batch:
        flash("No se encontró ese token.", "warning")
        return redirect(url_for("importar"))

    tab = (request.args.get("tab") or "ontracking").strip().lower()
    if tab not in ("ontracking", "log", "mapa"):
        tab = "ontracking"

    page = clamp_int(request.args.get("page"), default=1, lo=1, hi=10_000)
    per_page = clamp_int(request.args.get("per_page"), default=25, lo=10, hi=200)

    counts = {
        "ontracking": db_session.query(func.count(OntrackingRow.id)).filter_by(batch_id=batch.id).scalar() or 0,
        "log": db_session.query(func.count(CardLogRow.id)).filter_by(batch_id=batch.id).scalar() or 0,
        "mapa": db_session.query(func.count(RoomMapRow.id)).filter_by(batch_id=batch.id).scalar() or 0,
    }

    def page_query(model):
        q = db_session.query(model).filter_by(batch_id=batch.id).order_by(model.id.asc())
        total = counts[tab]
        offset = (page - 1) * per_page
        rows = q.offset(offset).limit(per_page).all()
        return total, rows

    columns_map = {
        "ontracking": ["MODULO", "LUGAR", "HABITACION", "EMPRESA", "ID", "CAMA", "INICIO", "TERMINO",
                       "DIA", "CAMAS_OCUPADAS", "RUT", "NOMBRE"],
        "log": ["NRO_TARJETA", "NRO_HABITACION", "HABITACION", "METODO_APERTURA_PUERTA",
                "TIPO_DE_TARJETA", "FECHA", "DUENO_CODIGO", "DUENO_NOMBRE"],
        "mapa": ["HABITACION", "MODULO", "PISO", "HKEYPLUS"],
    }

    model_map = {
        "ontracking": OntrackingRow,
        "log": CardLogRow,
        "mapa": RoomMapRow,
    }

    total, rows = page_query(model_map[tab])
    total_pages = max(1, (total + per_page - 1) // per_page)
    page = min(page, total_pages)

    # Convertir a lista raw para template
    raw_rows = [r.raw for r in rows]

    return render_template(
        "preview.html",
        token=token,
        created_at=batch.created_at,
        tab=tab,
        counts=counts,
        columns=columns_map[tab],
        rows=raw_rows,
        page=page,
        per_page=per_page,
        total=total,
        total_pages=total_pages,
    )


@app.get("/descargar/<token>/<dataset>.csv")
def descargar_csv(token: str, dataset: str):
    batch = db_session.query(UploadBatch).filter_by(token=token).first()
    if not batch:
        abort(404)

    model_map = {
        "ocupacion": OntrackingRow,
        "log_tarjetas": CardLogRow,
        "map_habitaciones": RoomMapRow,
    }
    if dataset not in model_map:
        abort(404)

    rows = (
        db_session.query(model_map[dataset])
        .filter_by(batch_id=batch.id)
        .order_by(model_map[dataset].id.asc())
        .all()
    )
    if not rows:
        abort(404)

    df = pd.DataFrame([r.raw for r in rows])
    bio = BytesIO()
    df.to_csv(bio, index=False, encoding="utf-8-sig")
    bio.seek(0)

    return send_file(
        bio,
        as_attachment=True,
        download_name=f"{dataset}_{token[:8]}.csv",
        mimetype="text/csv; charset=utf-8"
    )


# ----------------------------
# CONCILIACIÓN (por fecha usando mapa)
# ----------------------------
def compute_conciliacion(batch_id: int, target: date):
    date_str = target.isoformat()

    # 1) Mapa
    map_rows = (
        db_session.query(RoomMapRow)
        .filter_by(batch_id=batch_id)
        .order_by(RoomMapRow.id.asc())
        .all()
    )
    map_by_on = {}      # ontracking HABITACION -> dict
    map_by_hkey = {}    # hkeyplus -> ontracking HABITACION (para detectar logs sin mapa)
    for m in map_rows:
        on_h = (m.habitacion or "").strip()
        hk = (m.hkeyplus or "").strip()
        map_by_on[on_h] = {"on": on_h, "hkey": hk, "modulo": m.modulo or "", "piso": m.piso or ""}
        if hk:
            map_by_hkey[hk] = on_h

    # 2) Ontracking activos (fecha dentro de inicio/termino)
    # Como INICIO/TERMINO son "YYYY-MM-DD", comparación lexicográfica funciona.
    on_rows = (
        db_session.query(OntrackingRow)
        .filter_by(batch_id=batch_id)
        .filter(OntrackingRow.inicio != "")
        .filter(OntrackingRow.termino != "")
        .filter(OntrackingRow.inicio <= date_str)
        .filter(OntrackingRow.termino >= date_str)
        .all()
    )
    on_active = {}  # habitacion_on -> list[dict]
    for r in on_rows:
        h = (r.habitacion or "").strip()
        on_active.setdefault(h, []).append({
            "rut": r.rut or "",
            "nombre": r.nombre or "",
            "empresa": r.empresa or "",
            "cama": r.cama or "",
            "inicio": r.inicio or "",
            "termino": r.termino or "",
        })

    # 3) Logs del día
    log_rows = (
        db_session.query(CardLogRow)
        .filter_by(batch_id=batch_id)
        .order_by(CardLogRow.id.asc())
        .all()
    )
    log_by_hkey = {}  # log HABITACION (hkeyplus) -> list[dict]
    for r in log_rows:
        dt = parse_log_datetime(r.fecha)
        if not dt or dt.date() != target:
            continue
        hk = (r.habitacion or "").strip()
        log_by_hkey.setdefault(hk, []).append({
            "fecha": dt,
            "nro_tarjeta": r.nro_tarjeta or "",
            "dueno": (r.dueno_nombre or r.dueno_codigo or "").strip(),
            "metodo": (r.metodo_apertura_puerta or "").strip(),
        })

    # 4) Resultados
    results = []

    def summarize_people(lst):
        if not lst:
            return ""
        names = [x.get("nombre", "").strip() for x in lst if x.get("nombre")]
        names = [n for n in names if n]
        if not names:
            return ""
        if len(names) <= 2:
            return " · ".join(names)
        return f"{names[0]} · {names[1]} · +{len(names)-2}"

    # 4a) Habitaciones del mapa (principal)
    for on_h, info in map_by_on.items():
        hk = info["hkey"]
        people = on_active.get(on_h, [])
        logs = log_by_hkey.get(hk, [])

        ocupada = len(people) > 0
        log_n = len(logs)
        last_dt = max((x["fecha"] for x in logs), default=None)

        if ocupada and log_n > 0:
            status = "OK"
        elif ocupada and log_n == 0:
            status = "OCUPADA_SIN_LOG"
        elif (not ocupada) and log_n > 0:
            status = "LOG_SIN_OCUPACION"
        else:
            status = "SIN_MOVIMIENTOS"

        results.append({
            "habitacion_ontracking": on_h,
            "habitacion_log": hk,
            "modulo": info["modulo"],
            "piso": info["piso"],
            "ocupada": "SI" if ocupada else "NO",
            "ocupantes": summarize_people(people),
            "log_eventos": log_n,
            "log_ultimo": last_dt.strftime("%H:%M:%S") if last_dt else "",
            "status": status,
        })

    # 4b) Ontracking sin mapa
    for on_h, people in on_active.items():
        if on_h in map_by_on:
            continue
        results.append({
            "habitacion_ontracking": on_h,
            "habitacion_log": "",
            "modulo": "",
            "piso": "",
            "ocupada": "SI",
            "ocupantes": summarize_people(people),
            "log_eventos": 0,
            "log_ultimo": "",
            "status": "SIN_MAPA",
        })

    # 4c) Log sin mapa
    for hk, logs in log_by_hkey.items():
        if hk in map_by_hkey:
            continue
        last_dt = max((x["fecha"] for x in logs), default=None)
        results.append({
            "habitacion_ontracking": "",
            "habitacion_log": hk,
            "modulo": "",
            "piso": "",
            "ocupada": "NO",
            "ocupantes": "",
            "log_eventos": len(logs),
            "log_ultimo": last_dt.strftime("%H:%M:%S") if last_dt else "",
            "status": "SIN_MAPA",
        })

    # Orden: primero por ontracking, luego por log
    def key(r):
        return (r["habitacion_ontracking"] or "~~~~", r["habitacion_log"] or "~~~~")
    results.sort(key=key)

    # Conteos
    counts = {}
    for r in results:
        counts[r["status"]] = counts.get(r["status"], 0) + 1

    return results, counts


@app.get("/conciliacion")
def conciliacion():
    token = (request.args.get("token") or session.get("last_token") or "").strip()
    date_str = (request.args.get("date") or "").strip()  # YYYY-MM-DD
    page = clamp_int(request.args.get("page"), default=1, lo=1, hi=10_000)
    per_page = clamp_int(request.args.get("per_page"), default=25, lo=10, hi=200)

    if not token:
        return render_template("conciliacion.html", token="", date_str="", has_run=False)

    batch = db_session.query(UploadBatch).filter_by(token=token).first()
    if not batch:
        flash("Token no válido.", "danger")
        return redirect(url_for("importar"))

    if not date_str:
        return render_template("conciliacion.html", token=token, date_str="", has_run=False)

    try:
        target_date = datetime.strptime(date_str, "%Y-%m-%d").date()
    except Exception:
        flash("Fecha inválida.", "danger")
        return redirect(url_for("conciliacion", token=token))

    results, status_counts = compute_conciliacion(batch.id, target_date)

    total = len(results)
    total_pages = max(1, (total + per_page - 1) // per_page)
    page = min(page, total_pages)
    start = (page - 1) * per_page
    end = start + per_page
    page_rows = results[start:end]

    return render_template(
        "conciliacion.html",
        token=token,
        date_str=date_str,
        has_run=True,
        rows=page_rows,
        status_counts=status_counts,
        page=page,
        per_page=per_page,
        total=total,
        total_pages=total_pages,
    )


@app.get("/conciliacion/export.csv")
def conciliacion_export():
    token = (request.args.get("token") or "").strip()
    date_str = (request.args.get("date") or "").strip()
    if not token or not date_str:
        abort(404)

    batch = db_session.query(UploadBatch).filter_by(token=token).first()
    if not batch:
        abort(404)

    try:
        target_date = datetime.strptime(date_str, "%Y-%m-%d").date()
    except Exception:
        abort(404)

    results, _ = compute_conciliacion(batch.id, target_date)
    df = pd.DataFrame(results)

    bio = BytesIO()
    df.to_csv(bio, index=False, encoding="utf-8-sig")
    bio.seek(0)

    return send_file(
        bio,
        as_attachment=True,
        download_name=f"conciliacion_{token[:8]}_{date_str}.csv",
        mimetype="text/csv; charset=utf-8"
    )


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.getenv("PORT", "10000")), debug=True)
